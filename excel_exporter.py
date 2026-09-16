"""
excel_exporter.py
스크리닝된 종목 분석 데이터를 전문가용 엑셀(.xlsx) 및 CSV 형식으로
서식화하여 내보내는 모듈 (오름차순/내림차순 자동 필터 토글 지원)
"""

import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


EXPORT_HEADERS = [
    "종목명",
    "종목코드",
    "시장",
    "시가총액(억원)",
    "현재가",
    "거래량",
    "이동평균선 배열",
    "MACD",
    "RSI",
    "볼린저 밴드",
    "스토캐스틱",
    "거래량 수급",
    "종합 점수"
]


def prepare_export_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """
    원본 스크리닝 데이터프레임에서 사용자가 요청한 헤더 규격에 맞게 변환합니다.
    이전 세션 데이터나 누락된 컬럼에 대해 안전하게 fallback 처리합니다.
    """
    if df is None or df.empty:
        return pd.DataFrame(columns=EXPORT_HEADERS)

    df_out = pd.DataFrame()
    df_out["종목명"] = df.get("종목명", "")
    df_out["종목코드"] = df["종목코드"].astype(str).str.zfill(6) if "종목코드" in df else ""
    df_out["시장"] = df.get("시장", "")

    # 시가총액(억원)
    if "시가총액" in df.columns:
        df_out["시가총액(억원)"] = pd.to_numeric(df["시가총액"], errors="coerce").fillna(0).astype(int)
    elif "시가총액_원" in df.columns:
        df_out["시가총액(억원)"] = (pd.to_numeric(df["시가총액_원"], errors="coerce").fillna(0) / 100_000_000).round().astype(int)
    else:
        df_out["시가총액(억원)"] = 0

    df_out["현재가"] = pd.to_numeric(df.get("현재가", 0), errors="coerce").fillna(0).astype(int)

    if "거래량" in df.columns and pd.api.types.is_numeric_dtype(df["거래량"]):
        df_out["거래량"] = pd.to_numeric(df["거래량"], errors="coerce").fillna(0).astype(int)
    else:
        df_out["거래량"] = 0

    df_out["이동평균선 배열"] = df.get("이동평균선 배열", "")
    df_out["MACD"] = df.get("MACD", "")
    df_out["RSI"] = df.get("RSI", "")
    df_out["볼린저 밴드"] = df.get("볼린저 밴드", "")
    df_out["스토캐스틱"] = df.get("스토캐스틱", "")
    df_out["거래량 수급"] = df.get("거래량 수급", df.get("거래량", ""))
    df_out["종합 점수"] = pd.to_numeric(df.get("종합 점수", 0), errors="coerce").fillna(0).astype(int)

    return df_out


def create_excel_bytes(df: pd.DataFrame) -> bytes:
    """
    전문가용 서식(헤더 색상, 테두리, 너비 자동 조절, 점수 조건부 강조, 자동 필터 토글)이
    적용된 .xlsx 엑셀 바이너리를 생성합니다.
    """
    df_export = prepare_export_dataframe(df)

    wb = Workbook()
    ws = wb.active
    ws.title = "기술적 앙상블 스크리닝"

    # 스타일 정의
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # 짙은 슬레이트 네이비
    header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    
    data_font = Font(name="맑은 고딕", size=9)
    bold_data_font = Font(name="맑은 고딕", size=9, bold=True)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    buy_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # 부드러운 그린
    buy_font = Font(name="맑은 고딕", size=9, bold=True, color="166534")

    overweight_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")  # 부드러운 블루
    overweight_font = Font(name="맑은 고딕", size=9, bold=True, color="075985")

    # 1. 헤더 작성
    headers = list(df_export.columns)
    ws.append(headers)
    
    ws.row_dimensions[1].height = 26
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    # 2. 데이터 행 작성
    for row_idx, row_data in enumerate(df_export.itertuples(index=False), 2):
        ws.append(list(row_data))
        ws.row_dimensions[row_idx].height = 20

        score_val = row_data[-1]  # 종합 점수

        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border

            # 숫자 열 서식 적용 (#,##0)
            if col_name in ["시가총액(억원)", "현재가", "거래량"]:
                cell.alignment = align_right
                cell.number_format = "#,##0"
            elif col_name in ["종목코드", "시장"]:
                cell.alignment = align_center
            elif col_name in ["종목명"]:
                cell.font = bold_data_font
                cell.alignment = align_left
            elif col_name in ["종합 점수"]:
                cell.alignment = align_right
                cell.number_format = "+#,##0;-#,##0;0"
            else:
                cell.alignment = align_center

            # 종합 점수 하이라이트
            if col_name == "종합 점수":
                if score_val >= 45:
                    cell.fill = buy_fill
                    cell.font = buy_font
                elif score_val >= 15:
                    cell.fill = overweight_fill
                    cell.font = overweight_font

    # 3. 오름차순/내림차순 토글 필터(AutoFilter) 적용
    last_col_letter = get_column_letter(len(headers))
    last_row = len(df_export) + 1
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    # 4. 열 너비 자동 조정
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            str_len = len(val_str.encode('utf-8'))
            max_len = max(max_len, str_len)
        adjusted_width = max(10, int(max_len * 0.75) + 3)
        ws.column_dimensions[col_letter].width = min(adjusted_width, 32)

    # 5. 바이트 버퍼 저장
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def create_csv_bytes(df: pd.DataFrame) -> bytes:
    """
    호환성을 위한 UTF-8-SIG CSV 바이너리를 생성합니다.
    """
    df_export = prepare_export_dataframe(df)
    return df_export.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig")
